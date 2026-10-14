"""
data_loader.py — fetches EUR/USD candles from yfinance and saves xlsx to Euro_Datasets/

Usage:
    from data_loader import load_candles
    candles = load_candles("1min")

Supported intervals: 1min, 3min, 5min, 15min, 30min
"""
from datetime import datetime
from pathlib import Path
import yfinance as yf
import openpyxl

EURO_DIR = Path(__file__).parent / "Euro_Datasets"
EURO_DIR.mkdir(exist_ok=True)

FILE_MAP = {
    "1min":  "candlestick_eur_1min.xlsx",
    "3min":  "candlestick_eur_3min.xlsx",
    "5min":  "candlestick_eur_5min.xlsx",
    "15min": "candlestick_eur_15min.xlsx",
    "30min": "candlestick_eur_30min.xlsx",
}

YF_INTERVAL = {
    "1min":  "1m",
    "5min":  "5m",
    "15min": "15m",
    "30min": "30m",
}

YF_PERIOD = {
    "1min":  "7d",   # yfinance max for 1m
    "5min":  "60d",  # max for 5m
    "15min": "60d",
    "30min": "60d",
}

def _resample_3min(df):
    """yfinance has no 3min interval — resample from 1min."""
    df = df.resample("3min").agg({
        "Open": "first",
        "High": "max",
        "Low": "min",
        "Close": "last",
        "Volume": "sum",
    }).dropna()
    return df

def _fetch(interval):
    """Download from yfinance and return a DataFrame."""
    yf_interval = YF_INTERVAL.get(interval)
    if yf_interval:
        return yf.download("EURUSD=X", interval=yf_interval, period=YF_PERIOD[interval])
    # 3min: resample from 1m (7d max)
    df = yf.download("EURUSD=X", interval="1m", period="7d")
    if hasattr(df.columns, "get_level_values"):
        df.columns = df.columns.get_level_values(0)
    return _resample_3min(df)

def _save_xlsx(df, filename):
    """Save DataFrame to xlsx matching the existing column format."""
    # flatten multi-level columns (yfinance returns ("High", "EURUSD=X") style)
    if hasattr(df.columns, "get_level_values"):
        df.columns = df.columns.get_level_values(0)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candles"
    # header row matching existing datasets
    ws.append(["time", "open", "high", "low", "close", "vol", "x", "y", "datetime"])
    count = 0
    for idx, row in df.iterrows():
        ts = int(idx.timestamp()) if hasattr(idx, "timestamp") else 0
        o, h, l, c, v = float(row["Open"]), float(row["High"]), float(row["Low"]), float(row["Close"]), float(row["Volume"])
        dt = idx.strftime("%Y-%m-%d %H:%M:%S") if hasattr(idx, "strftime") else ""
        # skip flat candles (weekend / after-hours noise)
        if h == l == o == c:
            continue
        # swap H and L if inverted
        if h < l:
            h, l = l, h
        ws.append([ts, o, h, l, c, v or 0, 0, 0, dt])
        count += 1
    wb.save(filename)
    return count

def load_candles(interval, limit=6000):
    filename = FILE_MAP.get(interval)
    if not filename:
        raise ValueError(f"Unknown interval: {interval}. Use 1min, 3min, 5min, 15min, 30min")

    filepath = EURO_DIR / filename

    # always fetch fresh and overwrite the file
    print(f"Fetching {interval} data from yfinance...")
    df = _fetch(interval)
    # keep only the newest candles
    df = df.tail(limit)
    n = _save_xlsx(df, filepath)
    print(f"Saved {n} candles to {filepath}")
    
    # load from xlsx
    wb = openpyxl.load_workbook(str(filepath), read_only=True)
    ws = wb.active
    candles = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # skip header
        o, h, l, c = row[1], row[2], row[3], row[4]
        vol = row[5] or 0
        ts = row[0]
        candles.append((float(o), float(h), float(l), float(c), vol, datetime.fromtimestamp(ts)))
    wb.close()
    return candles
